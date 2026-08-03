#!/usr/bin/env python3
"""
Renderuje PPTX slajdove u PNG bez LibreOffice-a.
Parsira oblike preko python-pptx i kompozituje ih sa Pillow.
"""

import io
import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DECK = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else "native")
OUT.mkdir(parents=True, exist_ok=True)

RENDER_W = 2400  # px preko cele sirine slajda

SANS = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
SANS_B = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
EMOJI = "/usr/share/fonts/truetype/noto/NotoColorEmoji.ttf"
EMOJI_NATIVE = 109  # CBDT bitmap font se ucitava samo na ovoj velicini

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}

# Opseg NE sme da obuhvati strelice U+2190-21FF: tekstualni font ih ima, a
# emoji font nema, pa bi '→' bilo poslato u pogresan font i nestalo.
EMOJI_RE = re.compile(
    "([\U0001F000-\U0001FAFF☀-➿⬀-⯿️⌚-⌛⏩-⏳]+)"
)

_font_cache = {}


def font(size, bold=False):
    size = max(1, int(round(size)))
    key = (size, bold)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(SANS_B if bold else SANS, size)
    return _font_cache[key]


_emoji_font = None


def emoji_font():
    global _emoji_font
    if _emoji_font is None:
        try:
            _emoji_font = ImageFont.truetype(EMOJI, EMOJI_NATIVE)
        except OSError:
            _emoji_font = False
    return _emoji_font


# ---------------------------------------------------------------- utili

def solid_fill(shape):
    """Vraca (rgb, alpha) za solidFill ili None. python-pptx ne izlaze alfu."""
    el = shape._element
    for sp_pr in el.iter("{%s}spPr" % NS["a"]):
        fill = sp_pr.find("{%s}solidFill" % NS["a"])
        if fill is None:
            continue
        clr = fill.find("{%s}srgbClr" % NS["a"])
        if clr is None:
            return None
        rgb = tuple(int(clr.get("val")[i:i + 2], 16) for i in (0, 2, 4))
        alpha_el = clr.find("{%s}alpha" % NS["a"])
        alpha = int(alpha_el.get("val")) / 100000 if alpha_el is not None else 1.0
        return rgb, alpha
    return None


def run_color(run, default=(255, 255, 255)):
    try:
        c = run.font.color
        if c is not None and c.type is not None and c.rgb is not None:
            return tuple(c.rgb)
    except Exception:
        pass
    # fallback: direktno iz XML-a
    for clr in run._r.iter("{%s}srgbClr" % NS["a"]):
        v = clr.get("val")
        return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))
    return default


def measure(text, fnt, px):
    """Sirina teksta uz emoji koji se meri kao kvadrat visine px."""
    total = 0
    for part in EMOJI_RE.split(text):
        if not part:
            continue
        if EMOJI_RE.fullmatch(part):
            total += px * 1.15 * len(part)
        else:
            total += fnt.getlength(part)
    return total


def draw_rich(img, text, x, y, fnt, px, color):
    """Crta tekst; emoji segmenti idu preko Noto Color Emoji i skaliraju se."""
    ef = emoji_font()
    cx = x
    for part in EMOJI_RE.split(text):
        if not part:
            continue
        if EMOJI_RE.fullmatch(part) and ef:
            for ch in part:
                if ch in "️‍":
                    continue
                tile = Image.new("RGBA", (EMOJI_NATIVE + 40, EMOJI_NATIVE + 40), (0, 0, 0, 0))
                try:
                    ImageDraw.Draw(tile).text((0, 0), ch, font=ef, embedded_color=True)
                except Exception:
                    tile = None
                bbox = tile.getbbox() if tile else None
                if not bbox:
                    # Emoji font nema ovaj znak - crtamo ga tekstualnim fontom
                    # umesto da ga tiho ispustimo.
                    ImageDraw.Draw(img).text((cx, y), ch, font=fnt, fill=color + (255,))
                    cx += fnt.getlength(ch)
                    continue
                tile = tile.crop(bbox)
                target_h = int(px * 1.0)
                ratio = target_h / tile.height
                tile = tile.resize((max(1, int(tile.width * ratio)), target_h), Image.LANCZOS)
                img.alpha_composite(tile, (int(cx), int(y + px * 0.12)))
                cx += tile.width + px * 0.12
        else:
            ImageDraw.Draw(img).text((cx, y), part, font=fnt, fill=color + (255,))
            cx += fnt.getlength(part)


# ---------------------------------------------------------------- shapes

def render_picture(canvas, shape, S):
    try:
        blob = shape.image.blob
    except Exception:
        return
    try:
        pic = Image.open(io.BytesIO(blob)).convert("RGBA")
    except Exception:
        return

    # kropovanje definisano u srcRect
    cl, cr, ct, cb = (shape.crop_left, shape.crop_right,
                      shape.crop_top, shape.crop_bottom)
    if any(abs(v) > 1e-6 for v in (cl, cr, ct, cb)):
        w, h = pic.size
        box = (int(w * cl), int(h * ct), int(w * (1 - cr)), int(h * (1 - cb)))
        if box[2] > box[0] and box[3] > box[1]:
            pic = pic.crop(box)

    w = max(1, int(round(shape.width * S)))
    h = max(1, int(round(shape.height * S)))
    pic = pic.resize((w, h), Image.LANCZOS)

    if shape.rotation:
        pic = pic.rotate(-shape.rotation, expand=True, resample=Image.BICUBIC)

    canvas.alpha_composite(pic, (int(round(shape.left * S)), int(round(shape.top * S))))


def render_autoshape(canvas, shape, S):
    fill = solid_fill(shape)
    if not fill:
        return
    rgb, alpha = fill
    w = max(1, int(round(shape.width * S)))
    h = max(1, int(round(shape.height * S)))
    layer = Image.new("RGBA", (w, h), rgb + (int(alpha * 255),))
    canvas.alpha_composite(layer, (int(round(shape.left * S)), int(round(shape.top * S))))


def para_spacing(para, size_pt):
    """Vraca (prored_pt, razmak_pre_pt). Deck koristi spcPts, ne procente."""
    line_pt, before_pt = size_pt * 1.2, 0.0
    for tag, setter in (("lnSpc", "line"), ("spcBef", "before")):
        el = para._p.find(".//{%s}%s" % (NS["a"], tag))
        if el is None or not len(el):
            continue
        child = el[0]
        val = int(child.get("val"))
        if child.tag.endswith("spcPts"):
            v = val / 100
        else:  # spcPct - procenat velicine fonta
            v = size_pt * val / 100000
        if setter == "line":
            line_pt = v
        else:
            before_pt = v
    return line_pt, before_pt


def render_textbox(canvas, shape, S):
    tf = shape.text_frame
    if not tf.text.strip():
        return

    PT = S * 914400 / 72          # tacke -> pikseli
    box_w = shape.width * S
    lines = []                     # (text, font, px, color, align, line_px, before_px)

    for p_i, para in enumerate(tf.paragraphs):
        runs = [r for r in para.runs if r.text]
        if not runs:
            continue

        r0 = runs[0]
        pt = r0.font.size.pt if r0.font.size else 18
        px = pt * PT
        line_pt, before_pt = para_spacing(para, pt)
        line_px, before_px = line_pt * PT, (before_pt * PT if p_i else 0)

        bold = bool(r0.font.bold)
        fnt = font(px, bold)
        color = run_color(r0)
        text = "".join(r.text for r in runs)

        def wrap(f, size_px):
            out, current = [], []
            for wd in text.split():
                trial = " ".join(current + [wd])
                if measure(trial, f, size_px) > box_w and current:
                    out.append(" ".join(current))
                    current = [wd]
                else:
                    current.append(wd)
            out.append(" ".join(current))
            return out

        chunks = wrap(fnt, px)

        # Okvir ima zadatu visinu, pa zna koliko redova prima. Zamena fonta
        # (deck koristi uzi od Liberation Sans) ume da prelomi red koji je u
        # originalu stao u jedan - tada se font smanjuje tacno toliko da
        # prelom nestane, umesto da tekst iscuri preko sadrzaja ispod.
        allowed = max(1, round(shape.height * S / line_px)) if line_px else 1
        while len(chunks) > allowed and px > 8:
            px *= 0.96
            fnt = font(px, bold)
            chunks = wrap(fnt, px)

        for c_i, chunk in enumerate(chunks):
            lines.append((chunk, fnt, px, color, para.alignment,
                          line_px, before_px if c_i == 0 else 0))

    total_h = sum(lp + bp for *_, lp, bp in lines)
    top = shape.top * S
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        y = top + (shape.height * S - total_h) / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        y = top + shape.height * S - total_h
    else:
        y = top

    for text, fnt, px, color, align, line_px, before_px in lines:
        y += before_px
        if text:
            tw = measure(text, fnt, px)
            if align == PP_ALIGN.CENTER:
                x = shape.left * S + (box_w - tw) / 2
            elif align == PP_ALIGN.RIGHT:
                x = shape.left * S + box_w - tw
            else:
                x = shape.left * S
            # PIL crta od vrha ascendera; prored je mera osnovne linije
            draw_rich(canvas, text, x, y + (line_px - px) * 0.5, fnt, px, color)
        y += line_px


# ---------------------------------------------------------------- main

def main():
    prs = Presentation(DECK)
    S = RENDER_W / prs.slide_width
    H = int(round(prs.slide_height * S))
    print(f"Slajd {prs.slide_width}x{prs.slide_height} EMU -> {RENDER_W}x{H} px\n")

    for i, slide in enumerate(prs.slides, 1):
        canvas = Image.new("RGBA", (RENDER_W, H), (5, 8, 16, 255))
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            try:
                if shape.shape_type == 13:            # PICTURE
                    render_picture(canvas, shape, S)
                elif shape.has_text_frame:
                    if shape.shape_type == 1:          # AUTO_SHAPE sa tekstom
                        render_autoshape(canvas, shape, S)
                    render_textbox(canvas, shape, S)
                elif shape.shape_type == 1:
                    render_autoshape(canvas, shape, S)
            except Exception as e:
                print(f"   slajd {i}: preskocen oblik ({type(e).__name__}: {e})")

        path = OUT / f"native_{i:02d}.png"
        canvas.convert("RGB").save(path, "PNG")
        print(f"  {i:02d}/{len(prs.slides._sldIdLst)}  ->  {path.name}")


if __name__ == "__main__":
    main()
